#!/usr/bin/env python3
"""
Extract text, images, and speaker notes from PPTX files for frontend-slides.

The CLI prints JSON to stdout. Images are written to <output_dir>/assets/.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any


PICTURE_SHAPE_TYPE = 13
SUPPORTED_EXTENSIONS = {".pptx", ".pptm", ".potx"}


class CliError(Exception):
    """Expected CLI failure with a user-facing message."""


def _load_presentation_class():
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise CliError(
            "Missing dependency: python-pptx. Install it with "
            "`python3 -m pip install python-pptx`."
        ) from exc
    return Presentation


def _validate_input(path: Path) -> Path:
    path = path.expanduser().resolve()
    if not path.exists():
        raise CliError(f"Input file does not exist: {path}")
    if not path.is_file():
        raise CliError(f"Input path is not a file: {path}")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise CliError(
            f"Unsupported file type `{path.suffix}`. Supported types: {supported}. "
            "Legacy .ppt files must be converted to .pptx first."
        )
    return path


def _safe_stem(value: str, fallback: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip(".-")
    return cleaned or fallback


def _shape_text(shape: Any) -> str:
    text = getattr(shape, "text", "")
    return "\n".join(line.rstrip() for line in text.splitlines()).strip()


def extract_pptx(file_path: str | Path, output_dir: str | Path) -> dict[str, Any]:
    """Extract a PPTX file into JSON-friendly metadata and image assets."""
    input_path = _validate_input(Path(file_path))
    output_path = Path(output_dir).expanduser().resolve()
    assets_dir = output_path / "assets"

    Presentation = _load_presentation_class()
    try:
        prs = Presentation(str(input_path))
    except Exception as exc:
        raise CliError(f"Could not read PowerPoint file `{input_path}`: {exc}") from exc

    assets_dir.mkdir(parents=True, exist_ok=True)

    slides_data: list[dict[str, Any]] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        slide_data: dict[str, Any] = {
            "number": slide_index,
            "title": "",
            "content": [],
            "images": [],
            "notes": "",
        }

        title_shape = getattr(slide.shapes, "title", None)
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if getattr(shape, "has_text_frame", False):
                text = _shape_text(shape)
                if text:
                    if title_shape is not None and shape == title_shape:
                        slide_data["title"] = text
                    else:
                        slide_data["content"].append({
                            "type": "text",
                            "shape_index": shape_index,
                            "content": text,
                        })

            if getattr(shape, "shape_type", None) == PICTURE_SHAPE_TYPE:
                image = shape.image
                image_ext = image.ext.lower()
                image_name = (
                    f"slide{slide_index:02d}_img"
                    f"{len(slide_data['images']) + 1:02d}.{_safe_stem(image_ext, 'png')}"
                )
                image_path = assets_dir / image_name
                image_path.write_bytes(image.blob)
                slide_data["images"].append({
                    "path": f"assets/{image_name}",
                    "width_emu": int(shape.width),
                    "height_emu": int(shape.height),
                    "bytes": image_path.stat().st_size,
                })

        if slide.has_notes_slide:
            notes = _shape_text(slide.notes_slide.notes_text_frame)
            slide_data["notes"] = notes

        slides_data.append(slide_data)

    return {
        "source": str(input_path),
        "output_dir": str(output_path),
        "assets_dir": str(assets_dir),
        "slide_count": len(slides_data),
        "slides": slides_data,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Extract PPTX text, images, and notes into frontend-slides JSON."
    )
    parser.add_argument("input", help="Path to a .pptx/.pptm/.potx file.")
    parser.add_argument(
        "--out",
        required=True,
        help="Output directory. Extracted images are written to <out>/assets/.",
    )
    parser.add_argument(
        "--json-out",
        help="Optional path to also write the JSON metadata.",
    )
    parser.add_argument(
        "--pretty",
        action="store_true",
        help="Pretty-print JSON with indentation.",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    try:
        result = extract_pptx(args.input, args.out)
        json_text = json.dumps(
            result,
            ensure_ascii=False,
            indent=2 if args.pretty else None,
        )
        if args.json_out:
            json_path = Path(args.json_out).expanduser().resolve()
            json_path.parent.mkdir(parents=True, exist_ok=True)
            json_path.write_text(json_text + "\n", encoding="utf-8")
        print(json_text)
        return 0
    except CliError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print("error: interrupted", file=sys.stderr)
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
